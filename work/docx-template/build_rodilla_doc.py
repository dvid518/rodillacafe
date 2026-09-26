from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt

REFERENCE = Path(r"C:\Users\david\Downloads\David\cñz\hdp\Avance de Proyecto 1 - Grupo 3.docx")
OUTPUT = Path(r"C:\Users\david\Downloads\David\prog\rodilla-main\mdw\Avance de Proyecto 1 - Rodilla.docx")

doc = Document(REFERENCE)
body = doc._element.body
section_properties = body.sectPr
for child in list(body):
    if child is not section_properties:
        body.remove(child)

normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.0
normal.paragraph_format.space_before = Pt(5)
normal.paragraph_format.space_after = Pt(5)

def add_paragraph(text='', bold=False, size=None, align=None, before=None, after=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Arial'
    if size:
        run.font.size = Pt(size)
    return p

def add_heading(text, level=1):
    size = 14 if level == 1 else 12
    p = add_paragraph(text, bold=True, size=size, before=12 if level == 1 else 8, after=5)
    p.paragraph_format.keep_with_next = True
    return p

def add_body(text):
    p = add_paragraph(text)
    p.paragraph_format.first_line_indent = Pt(18)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def add_list(items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        ppr = p._p.get_or_add_pPr()
        numpr = OxmlElement('w:numPr')
        ilvl = OxmlElement('w:ilvl')
        ilvl.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '0')
        numid = OxmlElement('w:numId')
        numid.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '5')
        numpr.append(ilvl)
        numpr.append(numid)
        ppr.append(numpr)
        r = p.add_run(item)
        r.font.name = 'Arial'
        r.font.size = Pt(12)

# Title block
add_paragraph('UNIVERSIDAD TECNOLÓGICA DEL PERÚ', bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, before=55, after=8)
add_paragraph('Curso: Marcos de Desarrollo Web', size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=32)
add_paragraph('AVANCE DE PROYECTO FINAL 1', bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
add_paragraph('Sistema web para la gestión de la cafetería de especialidad Rodilla', bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=38)
add_paragraph('Integrantes: [Completar nombres de los integrantes]', size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
add_paragraph('Docente: [Completar nombre del docente]', size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
add_paragraph('Lima, 2026', size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=20)
doc.add_page_break()

add_paragraph('CAPÍTULO I: EL PROBLEMA', bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=15)

add_heading('1.1. Planteamiento del problema')
add_body('Rodilla es una cafetería de especialidad ficticia orientada a jóvenes entre 15 y 30 años. Su propuesta combina bebidas a base de café, alimentos complementarios y un espacio pensado para reuniones, estudio y actividades sociales. Debido a las expectativas de este público, la cafetería requiere canales digitales que permitan consultar información, elegir productos, realizar reservas y comunicarse con el negocio de manera rápida y ordenada.')
add_body('Actualmente, la información relacionada con los productos, pedidos, reservas y mensajes de los clientes puede administrarse de forma separada o mediante registros manuales. Esta situación dificulta que el personal consulte los datos actualizados, realice seguimiento de las solicitudes y tenga una visión integral de las operaciones. La ausencia de un espacio centralizado puede generar duplicidad de información, demoras en la atención, errores en el registro de reservas y una limitada trazabilidad de los pedidos.')
add_body('Asimismo, los clientes no cuentan con un canal integrado desde el cual puedan visualizar el menú, seleccionar productos, registrar un pedido, solicitar una reserva y enviar consultas. Esta dispersión afecta la experiencia del usuario y reduce la capacidad de la cafetería para organizar sus procesos de atención y toma de decisiones.')
add_body('Ante esta problemática, se propone un sistema web para Rodilla que centralice la información de productos, pedidos, reservas y mensajes de clientes. El sistema contará con interfaces para los clientes y un dashboard administrativo para gestionar los datos y eventos operativos de la cafetería. La propuesta incluye una base de datos PostgreSQL con datos simulados y considera la futura integración de un backend mediante una API.')

add_heading('1.2. Formulación del problema')
add_heading('1.2.1. Problema general', level=2)
add_body('¿De qué manera el desarrollo de un sistema web permitirá unificar la gestión de datos y eventos operativos de la cafetería de especialidad ficticia Rodilla?')
add_heading('1.2.2. Problemas específicos', level=2)
add_list([
    '¿Cómo se puede centralizar la información de productos, pedidos, reservas y mensajes de los clientes de Rodilla?',
    '¿De qué manera un dashboard administrativo puede facilitar el control de los datos registrados en la cafetería?',
    '¿Cómo puede un sistema web mejorar el proceso de consulta de productos, registro de pedidos y reservas de los clientes?',
    '¿De qué manera una base de datos relacional permitirá almacenar y organizar la información generada por las operaciones de Rodilla?'
])

add_heading('1.3. Justificación')
add_body('La propuesta se justifica en el aspecto práctico, tecnológico y académico. En el aspecto práctico, permitirá centralizar la información generada por la cafetería y facilitar el seguimiento de los productos, pedidos, reservas y mensajes. El dashboard administrativo reducirá la dependencia de registros dispersos y ofrecerá una vista organizada para el personal autorizado.')
add_body('En el aspecto tecnológico, el sistema contará con una interfaz web dirigida a los clientes y una zona de administración. Además, se plantea el diseño de una base de datos relacional en PostgreSQL, con tablas relacionadas, claves primarias, claves foráneas y restricciones que permitan mantener la consistencia de los datos. Esta base será el soporte para una posterior integración con un backend mediante servicios web.')
add_body('En el aspecto social y comercial, los clientes jóvenes dispondrán de un medio digital para consultar el menú, registrar pedidos, solicitar reservas y enviar mensajes. Esto facilitará la interacción con el negocio y permitirá que la cafetería responda a las necesidades de usuarios familiarizados con las plataformas digitales.')
add_body('Finalmente, en el aspecto académico, el proyecto permitirá aplicar conocimientos de desarrollo web, diseño de interfaces, modelado de bases de datos PostgreSQL y análisis de procesos de negocio, demostrando cómo una solución tecnológica puede responder a una problemática de gestión de información.')

add_heading('1.4. Delimitación del área de estudio')
add_body('El proyecto se desarrolla en las áreas de atención al cliente y administración de operaciones de la cafetería ficticia Rodilla. Los usuarios principales son los clientes, quienes podrán consultar el menú, seleccionar productos, registrar pedidos, solicitar reservas y enviar mensajes; y los administradores, quienes podrán gestionar productos, pedidos, reservas y mensajes desde un dashboard.')
add_body('El alcance del primer avance comprende el prototipo web funcional, el diseño e implementación de la base de datos PostgreSQL, la generación de datos simulados coherentes y la documentación de la arquitectura propuesta. El periodo de desarrollo corresponde al ciclo académico 2026-II.')

add_heading('1.5. Limitaciones de la investigación')
add_body('Rodilla es una empresa ficticia creada con fines académicos; por ello, no se dispone de registros históricos reales ni de información comercial obtenida de una cafetería en funcionamiento. Para esta fase se utilizarán datos simulados coherentes con el contexto del negocio.')
add_body('El tiempo establecido para el curso limita el alcance a las funciones esenciales: consulta de productos, pedidos, reservas, mensajes y administración de información. No se incluirán pasarelas de pago reales, facturación electrónica, integración con servicios de reparto, control avanzado de inventario ni conexiones con servicios externos.')
add_body('Durante este primer avance se presentará el diseño y la implementación de la base de datos, así como la definición de la arquitectura del backend. La programación e integración completa del backend con el frontend y PostgreSQL quedará prevista para una fase posterior.')

add_heading('1.6. Objetivos')
add_heading('1.6.1. Objetivo general', level=2)
add_body('Desarrollar un sistema web para la cafetería de especialidad ficticia Rodilla que permita unificar la gestión de productos, pedidos, reservas, mensajes y eventos operativos mediante un dashboard administrativo y una base de datos PostgreSQL.')
add_heading('1.6.2. Objetivos específicos', level=2)
add_list([
    'Diseñar una interfaz web que permita a los clientes consultar la información de Rodilla, visualizar el menú, registrar pedidos, realizar reservas y enviar mensajes.',
    'Diseñar una base de datos relacional en PostgreSQL para almacenar la información de usuarios, productos, categorías, pedidos, detalle de pedidos, reservas y mensajes.',
    'Generar datos simulados coherentes para las tablas transaccionales de la base de datos, garantizando registros suficientes para las pruebas del sistema.',
    'Implementar un dashboard administrativo que permita visualizar y gestionar los productos, pedidos, reservas y mensajes registrados.',
    'Definir la arquitectura de un backend que permita, en una fase posterior, conectar el frontend con la base de datos PostgreSQL mediante una API.'
])

add_heading('1.7. Propósito')
add_body('El propósito del proyecto es brindar a Rodilla una solución web que centralice la información operativa y mejore la interacción con sus clientes. La propuesta busca que los usuarios accedan de manera sencilla a los servicios de la cafetería, mientras que el personal administrativo disponga de un dashboard para organizar los datos y realizar un seguimiento de las operaciones.')
add_body('De esta manera, el sistema contribuirá a mejorar la organización de pedidos, reservas, productos y consultas de clientes, reduciendo la dispersión de información y estableciendo una base tecnológica escalable para futuras mejoras, como pagos digitales, control de inventario, reportes administrativos y servicios de entrega.')

doc.core_properties.title = 'Avance de Proyecto 1 - Rodilla'
doc.core_properties.subject = 'Capítulo I: El problema'
doc.core_properties.author = 'Grupo de proyecto'
doc.save(OUTPUT)
print(OUTPUT)
